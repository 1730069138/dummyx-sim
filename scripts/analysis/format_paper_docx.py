"""Apply submission-style formatting to a DOCX generated from the paper Markdown."""

from __future__ import annotations

import argparse
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_run_font(run, chinese: str = "宋体", latin: str = "Times New Roman", size: float = 10.5) -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), chinese)


def set_cell_margins(cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    """Use a Chinese academic three-line table: top, header rule, and bottom."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single" if edge in {"top", "bottom"} else "nil")
        node.set(qn("w:sz"), "12" if edge in {"top", "bottom"} else "0")
        node.set(qn("w:color"), "000000")

    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = tc_pr.first_child_found_in("w:tcBorders")
        if tc_borders is None:
            tc_borders = OxmlElement("w:tcBorders")
            tc_pr.append(tc_borders)
        bottom = tc_borders.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            tc_borders.append(bottom)
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:color"), "000000")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, end))


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


NUMERIC_CITATION_RE = re.compile(r"\[(?:\d+(?:\s*[-,]\s*\d+)*)\]")


def superscript_numeric_citations(paragraph) -> None:
    """Set bracketed numeric citations as superscript without lifting nearby prose."""
    for run in list(paragraph.runs):
        text = run.text
        matches = list(NUMERIC_CITATION_RE.finditer(text))
        if not matches:
            continue

        source_rpr = deepcopy(run._r.rPr)
        cursor = 0
        for match in matches:
            for segment, superscript in (
                (text[cursor:match.start()], False),
                (match.group(0), True),
            ):
                if not segment:
                    continue
                new_run = OxmlElement("w:r")
                if source_rpr is not None:
                    new_run.append(deepcopy(source_rpr))
                if superscript:
                    rpr = new_run.get_or_add_rPr()
                    vert_align = rpr.find(qn("w:vertAlign"))
                    if vert_align is None:
                        vert_align = OxmlElement("w:vertAlign")
                        rpr.append(vert_align)
                    vert_align.set(qn("w:val"), "superscript")
                text_node = OxmlElement("w:t")
                if segment[:1].isspace() or segment[-1:].isspace():
                    text_node.set(qn("xml:space"), "preserve")
                text_node.text = segment
                new_run.append(text_node)
                run._r.addprevious(new_run)
            cursor = match.end()

        if cursor < len(text):
            new_run = OxmlElement("w:r")
            if source_rpr is not None:
                new_run.append(deepcopy(source_rpr))
            text_node = OxmlElement("w:t")
            remainder = text[cursor:]
            if remainder[:1].isspace() or remainder[-1:].isspace():
                text_node.set(qn("xml:space"), "preserve")
            text_node.text = remainder
            new_run.append(text_node)
            run._r.addprevious(new_run)
        run._r.getparent().remove(run._r)


def split_table_with_continuation(doc, table_index: int, first_part_data_rows: int) -> None:
    """Split an appendix table and start its continuation on the next page."""
    table = doc.tables[table_index]
    table_xml = table._tbl
    rows = list(table_xml.findall(qn("w:tr")))
    split_at = first_part_data_rows + 1
    if split_at >= len(rows):
        return

    continuation_xml = deepcopy(table_xml)
    continuation_rows = list(continuation_xml.findall(qn("w:tr")))
    for row in rows[split_at:]:
        table_xml.remove(row)
    for row in continuation_rows[1:split_at]:
        continuation_xml.remove(row)

    caption = doc.add_paragraph("表 A1（续）")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = None
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.page_break_before = True
    for run in caption.runs:
        set_run_font(run, size=9)

    table_xml.addnext(caption._p)
    caption._p.addnext(continuation_xml)


EQUATION_LABELS = [
    "1",
    "2",
    "3a", "3b", "3c",
    "4a", "4b", "4c",
    "5a", "5b",
    "6", "7", "8", "9",
    "10a", "10b", "10c", "10d",
    "11", "12", "13", "14", "15", "16", "17", "18", "19", "20",
]


def number_display_equation(paragraph, number: str, center_twips: int, right_twips: int) -> None:
    """Keep the equation as editable OMML and place its number at the right margin."""
    math_paras = paragraph._p.xpath("./m:oMathPara")
    if not math_paras:
        return
    math_para = math_paras[0]
    equations = math_para.findall(qn("m:oMath"))
    if not equations:
        return
    equation = equations[0]
    math_para.remove(equation)
    paragraph._p.remove(math_para)

    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(center_twips / 20), WD_TAB_ALIGNMENT.CENTER)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(right_twips / 20), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run("\t")
    paragraph._p.append(equation)
    paragraph.add_run(f"\t({number})")


def number_inline_equation(paragraph, number: str, center_twips: int, right_twips: int) -> None:
    """Promote a Markdown inline-math equation line to centered editable OMML."""
    equations = paragraph._p.xpath("./m:oMath")
    if not equations:
        return
    equation = deepcopy(equations[0])
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(center_twips / 20), WD_TAB_ALIGNMENT.CENTER)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(right_twips / 20), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run("\t")
    paragraph._p.append(equation)
    paragraph.add_run(f"\t({number})")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    doc = Document(args.input)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(0)

    title_style = next((style for style in doc.styles if style.name == "Title"), None)
    if title_style is None:
        title_style = doc.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
    title = doc.paragraphs[0]
    title.style = title_style
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = None
    title.paragraph_format.space_after = Pt(12)
    for run in title.runs:
        set_run_font(run, chinese="黑体", size=15)
        run.bold = True

    for style_id, size in (("Heading1", 13), ("Heading2", 11.5)):
        style = next(style for style in doc.styles if style.style_id == style_id)
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.first_line_indent = None
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for paragraph in list(doc.paragraphs):
        if paragraph.style.name == "Image Caption":
            delete_paragraph(paragraph)

    reference_mode = False
    display_equation_index = 0
    for i, paragraph in enumerate(doc.paragraphs[1:], start=1):
        text = paragraph.text.strip()
        has_math = bool(paragraph._p.xpath(".//m:oMath"))
        has_drawing = bool(paragraph._p.xpath(".//w:drawing"))
        if text == "参考文献":
            reference_mode = True

        if i <= 4:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(3)
            size = 12 if i == 1 else (11 if i == 2 else 10.5)
            for run in paragraph.runs:
                set_run_font(run, chinese="黑体" if i == 1 else "宋体", size=size)
                run.bold = i == 1
        elif paragraph.style.name.startswith("Heading"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif has_drawing:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(2)
        elif re.match(r"^(图|表)\s*(?:\d+|A\d+)", text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(6)
            if text.startswith("表 A1"):
                paragraph.paragraph_format.page_break_before = True
            for run in paragraph.runs:
                set_run_font(run, size=9)
        elif has_math and re.fullmatch(r"\s*\(\d+[a-z]?\)\s*", text):
            if display_equation_index >= len(EQUATION_LABELS):
                raise ValueError("The document contains more numbered equations than expected.")
            number_inline_equation(
                paragraph,
                EQUATION_LABELS[display_equation_index],
                center_twips=4535,
                right_twips=9070,
            )
            display_equation_index += 1
        elif has_math and not text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)
            if paragraph._p.xpath("./m:oMathPara"):
                if display_equation_index >= len(EQUATION_LABELS):
                    raise ValueError("The document contains more numbered equations than expected.")
                number_display_equation(
                    paragraph,
                    EQUATION_LABELS[display_equation_index],
                    center_twips=4535,
                    right_twips=9070,
                )
                display_equation_index += 1
        elif reference_mode and re.match(r"^\[\d+\]", text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.left_indent = Cm(0.74)
            paragraph.paragraph_format.first_line_indent = Cm(-0.74)
            paragraph.paragraph_format.space_after = Pt(3)
        elif text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Pt(21)
            paragraph.paragraph_format.line_spacing = 1.25
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)

        if not paragraph.style.name.startswith("Heading"):
            for run in paragraph.runs:
                set_run_font(run)
        if not reference_mode:
            superscript_numeric_citations(paragraph)

    table_widths = (
        [0.8, 0.9, 1.2, 1.2, 3.5, 3.5, 2.2, 2.5],
        [3.0, 2.8, 2.8, 1.5, 1.7, 2.0, 2.0],
        [2.0, 3.7, 6.5, 3.6],
        [2.4, 4.4, 5.3, 3.7],
        [2.4, 4.4, 5.3, 3.7],
    )
    for idx, table in enumerate(doc.tables):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)
        widths = table_widths[idx] if idx < len(table_widths) else None
        if widths:
            for column, width in zip(table.columns, widths):
                column.width = Cm(width)
        for row_idx, row in enumerate(table.rows):
            if row_idx == 0:
                tr_pr = row._tr.get_or_add_trPr()
                repeat = OxmlElement("w:tblHeader")
                repeat.set(qn("w:val"), "true")
                tr_pr.append(repeat)
            for col_idx, cell in enumerate(row.cells):
                if widths and col_idx < len(widths):
                    cell.width = Cm(widths[col_idx])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.first_line_indent = None
                    paragraph.paragraph_format.space_before = Pt(1)
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.keep_with_next = False
                    paragraph.paragraph_format.keep_together = False
                    paragraph.paragraph_format.page_break_before = False
                    paragraph.paragraph_format.widow_control = False
                    for run in paragraph.runs:
                        set_run_font(run, size=8.0)
                        run.bold = row_idx == 0
                    superscript_numeric_citations(paragraph)

    split_table_with_continuation(doc, table_index=2, first_part_data_rows=14)

    footer = section.footer.paragraphs[0]
    footer.clear()
    add_page_number(footer)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
